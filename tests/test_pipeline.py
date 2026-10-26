from copy import deepcopy
from datetime import date
from decimal import Decimal
import json
import socket
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook

from office_synth.cli import main
from office_synth.config import initialize, load
from office_synth.generation import inspect_workbook
from office_synth.parsing import detect, parse
from office_synth.pipeline import Pipeline, models
from office_synth.planning import ordered, plan_errors
from office_synth.retrieval import DownloadError, resolve_public, acquire, fetch
from office_synth.schemas import FilePlan, EvidenceRecord, SeedBlueprint, WorldInstance
from office_synth.seeds import check_evidence, build_seed
from office_synth.storage import BudgetExceeded, Store, digest, read_json, write_json
from office_synth.taxonomy import sample, compatible
from office_synth.world import instantiate, world_errors


def test_offline_complete_and_export_isolation(built):
    p = built
    worlds = p.selected_worlds()
    assert len(worlds) == 2
    assert worlds[0].ledger() != worlds[1].ledger()
    assert worlds[0].group_key == worlds[1].group_key
    for w in worlds:
        workspace = p.workspace(w)
        assert read_json(workspace / "private/quality.json")["status"] == "accepted"
        files = list((workspace / "public").iterdir())
        assert len(files) >= 5
        assert {".docx", ".xlsx", ".pptx", ".pdf"} <= {f.suffix for f in files}
        ir = read_json(workspace / "private/ir/reconciliation_v1.json")
        assert ir["derivation"]["upstream_hash"] == digest((workspace / "public/03_settlement_data.xlsx").read_bytes())
    entries = p.export()
    assert len(entries) == 2
    for entry in entries:
        export = p.root / "export" / entry["workspace_id"]
        assert {f.name for f in export.iterdir()} == {"public", "manifest.json"}
        manifest = read_json(export / "manifest.json")
        assert "facts" not in manifest and "evidence" not in manifest
        assert all(f["path"].startswith("public/") for f in manifest["files"])
    assert p.store.counters() == {}


def test_deterministic_world_and_versions(built):
    p = built
    seed = models(p.root / "seeds.json", SeedBlueprint)[0]
    a, b = instantiate(seed, 0, p.config), instantiate(seed, 0, p.config)
    assert a.model_dump(exclude={"created_at"}) == b.model_dump(exclude={"created_at"})
    assert world_errors(a) == []
    assert a.rules[0]["valid_to"] == a.rules[1]["valid_from"]
    ledger = a.ledger()
    assert date.fromisoformat(ledger["due_date_v2"]) > date.fromisoformat(ledger["due_date_v1"])
    assert Decimal(ledger["payable_total"]) + Decimal(ledger["held_total"]) == Decimal(ledger["invoice_total"])


def test_tampered_amount_detected_and_local_repair(built):
    p = built
    w = p.selected_worlds()[0]
    workspace = p.workspace(w)
    untouched = workspace / "public/01_contract_v1.docx"
    original_mtime = untouched.stat().st_mtime_ns
    target = workspace / "public/04_reconciliation_v1.docx"
    doc = Document(target)
    for row in doc.tables[0].rows:
        if row.cells[0].text == "Payable total USD":
            row.cells[1].text = "999999.00"
    doc.save(target)
    p.validate()
    failures = [r for r in read_json(workspace / "private/validation.json") if r["status"] == "fail"]
    assert any(r["scope"] == "reconciliation_v1" and r["check_id"] == "facts" for r in failures)
    with pytest.raises(RuntimeError, match="Stale or invalid"):
        # Existing accepted quality cannot authorize export after a fresh tamper.
        write_json(workspace / "private/quality.json", {"status": "accepted", "quality_tier": "semantic"})
        p.export()
    p.repair()
    assert read_json(workspace / "private/quality.json")["status"] == "accepted"
    assert untouched.stat().st_mtime_ns == original_mtime
    assert list((workspace / "private/history/reconciliation_v1").iterdir())


def test_deleted_upstream_and_formula_tamper(built):
    p = built
    w = p.selected_worlds()[0]
    workspace = p.workspace(w)
    data = workspace / "public/03_settlement_data.xlsx"
    original_contract = (workspace / "public/01_contract_v1.docx").stat().st_mtime_ns
    wb = load_workbook(data)
    wb["Transactions"]["G2"] = "=1+1"
    wb.save(data)
    with pytest.raises(ValueError, match="formula"):
        inspect_workbook(data, w)
    p.validate()
    assert read_json(workspace / "private/quality.json")["status"] == "rejected"
    p.repair()
    assert read_json(workspace / "private/quality.json")["status"] == "accepted"
    assert (workspace / "public/01_contract_v1.docx").stat().st_mtime_ns == original_contract
    data.unlink()
    p.validate()
    results = read_json(workspace / "private/validation.json")
    assert any(r["check_id"] == "references" and r["status"] == "fail" for r in results)
    p.repair()
    assert data.exists()
    assert read_json(workspace / "private/quality.json")["status"] == "accepted"


def test_cycle_and_invalid_version_detection(built):
    p = built
    w = p.selected_worlds()[0]
    plans = models(p.workspace(w) / "private/plan.json", FilePlan)
    from office_synth.schemas import FileDependency
    next(x for x in plans if x.id == "data").dependencies.append(FileDependency(file_id="briefing", type="derived_from"))
    with pytest.raises(ValueError, match="cycle"):
        ordered(plans)
    assert any("cycle" in e for e in plan_errors(plans, w, p.config))
    plans = models(p.workspace(w) / "private/plan.json", FilePlan)
    plans[0].version_id = "v999"
    assert any("invalid version" in e for e in plan_errors(plans, w, p.config))


def test_resume_does_not_regenerate_or_call_services(built, monkeypatch):
    p = built
    before = {f: f.stat().st_mtime_ns for f in p.root.glob("workspaces/*/public/*")}
    monkeypatch.setattr("office_synth.pipeline.generate", lambda *a, **k: pytest.fail("Should reuse generated output"))
    p.run()
    assert all(path.stat().st_mtime_ns == mtime for path, mtime in before.items())
    assert p.reused >= 14
    assert p.store.counters() == {}


def test_budget_persistence_and_interruption_recovery(pipeline):
    p = pipeline
    p.config["budget"]["max_search_requests"] = 1
    p.budget.request("search_requests")
    with pytest.raises(BudgetExceeded):
        p.budget.request("search_requests")
    assert p.store.counters()["search_requests"] == 1
    path = p.root / "example.json"
    def interrupted():
        raise KeyboardInterrupt()
    with pytest.raises(KeyboardInterrupt):
        p.store.execute("demo", "object", {}, interrupted)
    def recovered():
        write_json(path, {"done": True})
        return ["example.json"]
    p.store.execute("demo", "object", {}, recovered)
    assert p.store.get_stage("demo", "object").attempts == 2
    assert p.store.get_stage("demo", "object").status.value == "succeeded"


def test_compatibility_and_bounded_sampling(config):
    assert not compatible({"role": "finance", "scenario": "bad"}, [{"effect": "deny", "when": {"role": ["finance"]}, "require": {"scenario": ["bad"]}}])
    a, coverage = sample(config["_taxonomy"], config)
    b, _ = sample(config["_taxonomy"], config)
    assert [c.id for c in a] == [c.id for c in b]
    assert {c.role_id for c in a} == {"finance", "procurement", "operations"}
    assert len(a) <= config["sampling"]["candidate_limit"]
    assert all(v > 0 for v in coverage["marginal"]["role"].values())


def test_evidence_position_and_insufficient_sources(built):
    from office_synth.schemas import SourceRecord
    p = built
    sources = models(p.root / "sources.json", SourceRecord)
    evidence = models(p.root / "evidence.json", EvidenceRecord)
    parsed = {s.id: read_json(p.root / s.parsed_path) for s in sources}
    assert check_evidence(evidence, sources, parsed) == []
    evidence[0].excerpt = "Unsupported claim about a nonexistent institution."
    assert check_evidence(evidence, sources, parsed)
    from office_synth.schemas import CandidateSeed
    candidate = models(p.root / "candidates.json", CandidateSeed)[0]
    seed = build_seed(candidate, sources[:1], evidence, p.config)
    assert seed.review_status == "needs_evidence"


@pytest.mark.parametrize("url", ["file:///etc/passwd", "http://127.0.0.1/a", "http://169.254.169.254/latest/meta-data", "http://[::1]/", "http://user:password@example.com"])
def test_ssrf_rejected(url):
    with pytest.raises(DownloadError):
        resolve_public(url, [])


def test_mixed_dns_rejected(monkeypatch):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(2, 1, 6, "", ("8.8.8.8", 80)), (2, 1, 6, "", ("10.0.0.1", 80))])
    with pytest.raises(DownloadError):
        resolve_public("http://example.com", [])


def test_login_mime_and_archive_rejection(pipeline, monkeypatch, tmp_path):
    monkeypatch.setattr("office_synth.retrieval.fetch", lambda *a: (b'<html><input type="password"></html>', "https://example.org/test.xlsx", "application/vnd.ms-excel", 200))
    source = acquire("https://example.org/test.xlsx", pipeline)
    assert source.download_status == "succeeded" and source.parse_status == "failed"
    assert (pipeline.root / source.local_path).exists()
    path = tmp_path / "bad.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("../word/document.xml", "x")
    with pytest.raises(ValueError, match="Unsafe ZIP"):
        detect(path.read_bytes())
    with pytest.raises(ValueError, match="Legacy"):
        detect(bytes.fromhex("D0CF11E0") + b"data")


def test_dry_run_no_writes_and_init_no_overwrite(tmp_path):
    config_dir = tmp_path / "configs"
    initialize(config_dir)
    path = config_dir / "pipeline.yaml"
    cfg = load(path)
    assert main(["run", "--config", str(path), "--dry-run"]) == 0
    from pathlib import Path
    assert not Path(cfg["_root"]).exists()
    original = path.read_bytes()
    with pytest.raises(FileExistsError):
        initialize(config_dir)
    assert path.read_bytes() == original


def test_native_format_source_parsing(built):
    for path in (built.workspace(built.selected_worlds()[0]) / "public").iterdir():
        fmt = detect(path.read_bytes(), name=path.name)
        result = parse(path, fmt)
        assert result["raw_hash"] == digest(path.read_bytes())
        assert result["units"]
        assert all(u["location"] for u in result["units"])


def test_unsupported_required_check_blocks_acceptance(built):
    built.config["validation"]["required_checks"].append("external_recalculation")
    built.validate()
    for w in built.selected_worlds():
        quality = read_json(built.workspace(w) / "private/quality.json")
        assert quality["status"] == "rejected" and quality["skipped_required"]


def test_three_roles_and_optional_communication(pipeline):
    p = pipeline
    p.config["sampling"]["target_seeds"] = 3
    p.config["generation"]["optional_formats"] = ["md", "eml"]
    p.run()
    assert {w.role for w in p.selected_worlds()} == {"procurement", "finance", "operations"}
    assert len(p.selected_worlds()) == 6
    assert all(read_json(p.workspace(w) / "private/quality.json")["status"] == "accepted" for w in p.selected_worlds())
    from email import policy
    from email.parser import BytesParser
    from email.utils import parsedate_to_datetime
    for world in p.selected_worlds():
        message = BytesParser(policy=policy.default).parsebytes((p.workspace(world) / "public/09_communication_eml.eml").read_bytes())
        assert str(parsedate_to_datetime(message["Date"]).date()) == world.ledger()["invoice_date"]


def test_brave_cache_and_auth_failure_without_retry(pipeline, monkeypatch):
    from office_synth.providers import BraveSearch, ProviderError
    import httpx
    monkeypatch.setenv("BRAVE_SEARCH_API_KEY", "test-secret")
    calls = []
    def response(*a, **k):
        calls.append(1)
        return httpx.Response(200, json={"web": {"results": [{"url": "https://example.org", "title": "Example", "description": "Selection only"}]}})
    monkeypatch.setattr(httpx, "get", response)
    search = BraveSearch(pipeline)
    assert search.search("query", {}, 1) == search.search("query", {}, 1)
    assert len(calls) == 1
    monkeypatch.setattr(httpx, "get", lambda *a, **k: httpx.Response(401, text="test-secret"))
    with pytest.raises(ProviderError, match="401") as exc:
        search.search("different query", {}, 1)
    assert "test-secret" not in str(exc.value)
    assert pipeline.store.counters()["search_requests"] == 2
