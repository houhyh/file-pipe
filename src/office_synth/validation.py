from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from .generation import inspect_workbook, LABELS
from .parsing import detect, parse
from .planning import plan_errors
from .schemas import FilePlan, ValidationResult, WorldInstance, money
from .storage import digest, read_json, safe_path, write_json
from .world import world_errors


def validate(world: WorldInstance, plans: list[FilePlan], workspace: Path, config: dict) -> list[ValidationResult]:
    results = []

    def check(scope, key, status, expected=None, observed=None, severity="blocking", location=None):
        results.append(ValidationResult(id=f"{scope}:{key}", scope=scope, check_id=key, status=status,
                                        expected=expected, observed=observed, severity=severity,
                                        location=location or {}, failure_class=key if status == "fail" else None,
                                        repair_hint="Repair affected file and regenerate dependent files" if status == "fail" else None))

    try:
        errors = world_errors(world)
    except (ValueError, KeyError, TypeError, ArithmeticError) as exc:
        errors = [f"Invalid fact structure: {exc}"]
    check(world.id, "world", "fail" if errors else "pass", "business invariants hold", errors)
    errors = plan_errors(plans, world, config)
    check(world.id, "dependencies", "fail" if errors else "pass", "valid plan and generation DAG", errors)
    public = workspace / "public"
    ledger = world.ledger()
    for plan in plans:
        path = safe_path(public, plan.relative_path)
        try:
            detected = detect(path.read_bytes(), "text/plain" if plan.format in {"md", "eml"} else "", path.name)
            if detected != plan.format and not (plan.format == "eml" and detected == "txt"):
                raise ValueError(f"Expected {plan.format}, detected {detected}")
            parsed = parse(path, detected)
            check(plan.id, "structure", "pass", "readable native content", {"units": len(parsed["units"]), "hash": digest(path.read_bytes())})
        except Exception as exc:
            check(plan.id, "structure", "fail", "readable native content", str(exc))
            continue
        missing_refs = [d.file_id for d in plan.dependencies if not any(p.id == d.file_id and safe_path(public, p.relative_path).is_file() for p in plans)]
        # Check actual rendered references as well as planned references.
        missing_refs += [p.relative_path for d in plan.dependencies for p in plans if p.id == d.file_id and p.relative_path not in parsed["text"]]
        check(plan.id, "references", "fail" if missing_refs else "pass", "references exist and appear in content", missing_refs)
        rules = [r for r in world.rules if r["version"] == plan.version_id]
        version_ok = len(rules) == 1 and rules[0]["valid_from"] == plan.valid_from and rules[0]["valid_to"] == plan.valid_to
        if plan.format != "xlsx":
            version_ok = version_ok and plan.version_id in parsed["text"] and plan.valid_from in parsed["text"] and plan.business_date in parsed["text"]
        check(plan.id, "versions", "pass" if version_ok else "fail", "version and validity match applicable rule", plan.version_id)
        try:
            if plan.format == "xlsx":
                derivation = inspect_workbook(path, world)
                check(plan.id, "calculations", "pass", "Decimal evaluator and formula caches agree with ledger", derivation)
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True, data_only=False)
                try:
                    terms = dict(wb["Terms"].iter_rows(min_row=2, max_row=8, max_col=2, values_only=True))
                    mismatches = [key for key in ["buyer", "supplier", "order_id", "invoice_id", "invoice_date", "payment_days_v1", "due_date_v1"] if str(terms.get(LABELS[key])) != str(ledger[key])]
                    for ri, line in enumerate(world.transactions, 2):
                        observed = [wb["Transactions"].cell(ri, ci).value for ci in range(1, 7)]
                        expected = [line["line_id"], line["product"], line["ordered"], line["received"], line["invoiced"]]
                        if observed[:5] != expected or money(observed[5]) != money(line["unit_price"]):
                            mismatches.append(f"Transactions!A{ri}:F{ri}")
                finally:
                    wb.close()
                if mismatches:
                    raise ValueError(f"Workbook facts mismatch: {mismatches}")
            else:
                # Check actual table/text values, never private IR alone.
                if plan.format == "docx":
                    from docx import Document
                    doc = Document(path)
                    fields = {r.cells[0].text: r.cells[1].text for t in doc.tables[:1] for r in t.rows[1:]}
                    mismatches = [fid for fid in plan.fact_ids if not fid.startswith("line_") and fields.get(LABELS.get(fid, fid)) != str(ledger[fid])]
                    if plan.role == "purchase_order":
                        expected_lines = [[str(i[k]) for k in ["product", "ordered", "received", "unit_price", "payable_amount"]] for i in world.transactions]
                        actual_lines = [[c.text for c in row.cells] for row in doc.tables[1].rows[1:]]
                        if actual_lines != expected_lines:
                            mismatches.append("transaction table")
                elif plan.format == "pptx":
                    from pptx import Presentation
                    tables = [shape.table for s in Presentation(path).slides for shape in s.shapes if shape.has_table]
                    fields = {r.cells[0].text: r.cells[1].text for t in tables for r in list(t.rows)[1:]}
                    mismatches = [fid for fid in plan.fact_ids if not fid.startswith("line_") and fields.get(LABELS.get(fid, fid)) != str(ledger[fid])]
                else:
                    mismatches = [fid for fid in plan.fact_ids if not fid.startswith("line_") and str(ledger[fid]) not in parsed["text"]]
                if mismatches:
                    raise ValueError(f"Rendered facts mismatch: {mismatches}")
                if any(d.type in {"derived_from", "summarizes"} for d in plan.dependencies):
                    data_plan = next(p for p in plans if p.id == "data")
                    derivation = inspect_workbook(public / data_plan.relative_path, world)
                    ir = read_json(workspace / "private" / "ir" / f"{plan.id}.json")
                    if ir.get("derivation", {}).get("upstream_hash") != derivation["upstream_hash"]:
                        raise ValueError("Derived content has a stale upstream hash")
                    check(plan.id, "calculations", "pass", "actual upstream read and computed result", derivation)
            check(plan.id, "facts", "pass", "rendered facts agree with ledger")
        except Exception as exc:
            check(plan.id, "facts", "fail", "rendered facts agree with ledger", str(exc))
            if not any(r.scope == plan.id and r.check_id == "calculations" for r in results):
                check(plan.id, "calculations", "fail", "valid upstream computations", str(exc))
        if plan.format == "pptx":
            from pptx import Presentation
            prs = Presentation(path)
            outside = [f"slide {si} shape {s.shape_id}" for si, slide in enumerate(prs.slides, 1) for s in slide.shapes
                       if s.left < 0 or s.top < 0 or s.left+s.width > prs.slide_width or s.top+s.height > prs.slide_height]
            check(plan.id, "geometry", "fail" if outside else "pass", "shapes inside slide", outside)
        if plan.format == "xlsx":
            external_recalc(path, plan, world, config, check)
        if plan.format in {"docx", "pptx", "pdf"}:
            visual(path, plan, workspace, config, check)
    return results


def libreoffice_convert(path: Path, fmt: str, output: Path):
    tool = shutil.which("libreoffice")
    if not tool:
        raise FileNotFoundError("LibreOffice is not installed")
    with tempfile.TemporaryDirectory(prefix="office-profile-") as profile:
        result = subprocess.run([tool, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", fmt,
                                 "--outdir", str(output), str(path.resolve())], capture_output=True, text=True, timeout=45)
    target = output / (path.stem + "." + fmt.split(":")[0])
    if result.returncode or not target.is_file():
        raise RuntimeError(f"LibreOffice conversion failed ({result.returncode}): {result.stderr[:200]}")
    return target


def external_recalc(path, plan, world, config, check):
    if not config["validation"]["recalculate"]:
        check(plan.id, "external_recalculation", "skipped", "optional external engine", "disabled", "warning")
        return
    try:
        with tempfile.TemporaryDirectory(prefix="office-recalc-") as temp:
            target = libreoffice_convert(path, "xlsx", Path(temp))
            from openpyxl import load_workbook
            wb = load_workbook(target, data_only=True, read_only=True)
            try:
                values = {key: wb["Summary"][cell].value for cell, key in [("B2", "payable_total"), ("B3", "invoice_total"), ("B4", "held_total")]}
                if any(money(value) != money(world.ledger()[key]) for key, value in values.items()):
                    raise ValueError("External recalculation differs from ledger")
            finally:
                wb.close()
            check(plan.id, "external_recalculation", "pass", "LibreOffice recalculated totals", values)
    except Exception as exc:
        # A real computation mismatch blocks acceptance; a missing/unusable engine is skipped.
        check(plan.id, "external_recalculation", "fail" if isinstance(exc, ValueError) else "skipped", "LibreOffice recalculation", str(exc), "blocking" if isinstance(exc, ValueError) else "warning")


def visual(path, plan, workspace, config, check):
    if not config["validation"]["render"]:
        check(plan.id, "visual", "skipped", "optional rendered page inspection", "disabled", "warning")
        return
    try:
        import importlib.util
        if importlib.util.find_spec("pymupdf") is None:
            raise FileNotFoundError("Optional PyMuPDF rasterizer is unavailable")
        import pymupdf
        output = workspace / "private" / "rendered" / plan.id
        output.mkdir(parents=True, exist_ok=True)
        pdf = path if plan.format == "pdf" else libreoffice_convert(path, "pdf", output)
        doc = pymupdf.open(pdf)
        problems = []
        for pi, page in enumerate(doc):
            if not page.get_text().strip():
                problems.append(f"blank or textless page {pi+1}")
            for block in page.get_text("blocks"):
                if block[0] < -1 or block[1] < -1 or block[2] > page.rect.width+1 or block[3] > page.rect.height+1:
                    problems.append(f"text beyond page {pi+1}")
            page.get_pixmap(matrix=pymupdf.Matrix(1, 1)).save(output / f"page-{pi+1}.png")
        pages = len(doc)
        doc.close()
        if not pages:
            problems.append("no pages")
        check(plan.id, "visual", "fail" if problems else "pass", "nonblank rendered pages, text within page bounds", {"pages": pages, "problems": problems})
    except Exception as exc:
        check(plan.id, "visual", "skipped", "rendered page inspection", str(exc), "warning")


def assessment(results: list[ValidationResult], config: dict) -> dict:
    required = set(config["validation"]["required_checks"])
    failed = [r.id for r in results if r.status == "fail" and r.severity == "blocking"]
    missing = [key for key in required if not any(r.check_id == key for r in results)]
    skipped_required = [r.id for r in results if r.check_id in required and r.status == "skipped"]
    if failed or missing or skipped_required:
        return {"status": "rejected", "quality_tier": "unaccepted", "failed": failed, "missing": missing, "skipped_required": skipped_required}
    visuals = [r for r in results if r.check_id == "visual"]
    tier = "rendered" if visuals and all(r.status == "pass" for r in visuals) else "semantic"
    return {"status": "accepted", "quality_tier": tier, "failed": [], "missing": [], "skipped_required": []}
