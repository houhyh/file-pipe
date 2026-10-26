from __future__ import annotations

import os
import tempfile
from decimal import Decimal
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import format_datetime
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from .schemas import FilePlan, WorldInstance, money
from .storage import atomic_bytes, digest, safe_path, write_json


LABELS = {"buyer": "Buyer", "supplier": "Supplier", "order_id": "Purchase order", "invoice_id": "Invoice",
          "order_date": "Order date", "receipt_date": "Receipt date", "invoice_date": "Invoice date",
          "invoice_total": "Invoice total USD", "payable_total": "Payable total USD", "held_total": "Held pending receipt USD",
          "payment_days_v1": "Payment days (original)", "due_date_v1": "Due date (original)",
          "payment_days_v2": "Payment days (amended)", "due_date_v2": "Due date (amended)"}
TITLES = {"contract": "Supply and Payment Agreement", "purchase_order": "Purchase Order",
          "reconciliation": "Supplier Settlement Reconciliation", "management_briefing": "Settlement Decision Briefing",
          "statement": "Supplier Settlement Statement", "amendment": "Payment Terms Amendment", "communication": "Settlement Update"}


def inspect_workbook(path: Path, world: WorldInstance) -> dict:
    """Compute from real workbook inputs; never fall back to ledger when parsing fails."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        ws = wb["Transactions"]
        lines = list(ws.iter_rows(min_row=2, max_row=4, values_only=True))
        payable, invoiced = Decimal("0"), Decimal("0")
        for i, row in enumerate(lines, 2):
            if row[6] != f"=ROUND(D{i}*F{i},2)" or row[7] != f"=ROUND(E{i}*F{i},2)":
                raise ValueError(f"Unexpected transaction formula at row {i}")
            if not 0 <= int(row[3]) <= int(row[2]) or int(row[4]) != int(row[2]):
                raise ValueError(f"Invalid workbook quantity at row {i}")
            payable += money(Decimal(str(row[3])) * Decimal(str(row[5])))
            invoiced += money(Decimal(str(row[4])) * Decimal(str(row[5])))
        expected_summary = {"B2": "=SUM(Transactions!G2:G4)", "B3": "=SUM(Transactions!H2:H4)", "B4": "=B3-B2"}
        for cell, formula in expected_summary.items():
            if wb["Summary"][cell].value != formula:
                raise ValueError(f"Unexpected summary formula at {cell}")
        result = {"payable_total": str(money(payable)), "invoice_total": str(money(invoiced)), "held_total": str(money(invoiced-payable))}
        for key, value in result.items():
            if money(value) != money(world.ledger()[key]):
                raise ValueError(f"Workbook computed {key}={value}, ledger={world.ledger()[key]}")
    finally:
        wb.close()
    cached = load_workbook(path, data_only=True, read_only=True)
    try:
        for cell, key in [("B2", "payable_total"), ("B3", "invoice_total"), ("B4", "held_total")]:
            if cached["Summary"][cell].value is None or money(cached["Summary"][cell].value) != money(result[key]):
                raise ValueError(f"Formula cache mismatch at Summary!{cell}")
    finally:
        cached.close()
    return {"method": "decimal_evaluation_of_allowlisted_formulas", "values": result,
            "upstream_hash": digest(path.read_bytes()), "external_engine_recalculated": False}


def content_ir(plan: FilePlan, world: WorldInstance, plans: list[FilePlan], public: Path, narrative: str | None = None) -> dict:
    values = world.ledger().copy()
    derivation = None
    if any(d.type in {"derived_from", "summarizes"} for d in plan.dependencies):
        upstream = next(p for p in plans if p.id == "data")
        derivation = inspect_workbook(safe_path(public, upstream.relative_path), world)
        values.update(derivation["values"])
    references = [next(p.relative_path for p in plans if p.id == d.file_id) for d in plan.dependencies]
    role_text = {
        "procurement": "Confirm supplier agreement and match purchase order, received quantities and invoice before releasing payment.",
        "finance": "Approve only the reconciled payable amount. Keep the unmatched balance on hold until the receipt exception is resolved.",
        "operations": "Check warehouse receipt quantities, notify the supplier of short delivery and coordinate resolution before settlement.",
    }.get(world.role, "Review and reconcile the settlement records.")
    paragraphs = ["Synthetic business example. All organizations and transactions are fictional.", role_text,
                  "Settlement uses received quantities multiplied by agreed unit prices. Unreceived items remain on hold."]
    if plan.role == "amendment":
        paragraphs += [f"By mutual agreement, the outstanding unpaid invoice moves from {values['payment_days_v1']} to {values['payment_days_v2']} payment days effective {plan.valid_from}.",
                       "The original agreement remains a historical record; this amendment applies to the identified unpaid invoice."]
    if narrative:
        paragraphs.append(narrative)
    return {"title": TITLES.get(plan.role, plan.role), "version": plan.version_id,
            "validity": f"{plan.version_id} | Issued {plan.business_date} | Effective {plan.valid_from}; ends {plan.valid_to or 'open'} (end exclusive)",
            "paragraphs": paragraphs,
            "facts": [{"id": fid, "label": LABELS.get(fid, fid), "value": str(values[fid])} for fid in plan.fact_ids if not fid.startswith("line_")],
            "lines": world.transactions if plan.role in {"purchase_order", "reconciliation", "statement"} else [],
            "references": references, "derivation": derivation}


def render_xlsx(path: Path, world: WorldInstance, references: list[str]):
    import xlsxwriter
    wb = xlsxwriter.Workbook(path)
    wb.set_calc_mode("auto")
    wb.set_properties({"title": "Fictional supplier settlement", "comments": "Formula caches computed with Decimal; Excel recalculates on open."})
    heading = wb.add_format({"bold": True, "bg_color": "#14334D", "font_color": "white", "text_wrap": True})
    currency = wb.add_format({"num_format": '"$"#,##0.00'})
    ws = wb.add_worksheet("Transactions")
    headers = ["Line", "Product", "Ordered", "Received", "Invoiced", "Unit price USD", "Payable USD", "Invoice USD"]
    ws.write_row(0, 0, headers, heading)
    ws.set_column(0, 0, 12)
    ws.set_column(1, 1, 25)
    ws.set_column(2, 7, 17)
    ws.freeze_panes(1, 2)
    for row, line in enumerate(world.transactions, 1):
        ws.write_row(row, 0, [line["line_id"], line["product"], line["ordered"], line["received"], line["invoiced"]])
        ws.write_number(row, 5, float(money(line["unit_price"])), currency)
        excel_row = row + 1
        ws.write_formula(row, 6, f"=ROUND(D{excel_row}*F{excel_row},2)", currency, float(money(line["payable_amount"])))
        ws.write_formula(row, 7, f"=ROUND(E{excel_row}*F{excel_row},2)", currency, float(money(line["ordered_amount"])))
    ws.autofilter(0, 0, 3, 7)
    ws.set_landscape()
    ws.fit_to_pages(1, 1)
    summary = wb.add_worksheet("Summary")
    summary.set_column(0, 0, 28)
    summary.set_column(1, 1, 25)
    summary.write_row(0, 0, ["Metric", "Value USD"], heading)
    ledger = world.ledger()
    for row, (name, formula, key) in enumerate([
        ("Payable", "=SUM(Transactions!G2:G4)", "payable_total"),
        ("Invoiced", "=SUM(Transactions!H2:H4)", "invoice_total"),
        ("Held", "=B3-B2", "held_total")], 1):
        summary.write(row, 0, name)
        summary.write_formula(row, 1, formula, currency, float(money(ledger[key])))
    chart = wb.add_chart({"type": "column"})
    chart.add_series({"name": "Settlement USD", "categories": "=Summary!$A$2:$A$4", "values": "=Summary!$B$2:$B$4"})
    chart.set_title({"name": "Invoice disposition"})
    chart.set_legend({"none": True})
    summary.insert_chart("D2", chart)
    terms = wb.add_worksheet("Terms")
    terms.set_column(0, 0, 30)
    terms.set_column(1, 1, 55)
    terms.write_row(0, 0, ["Field", "Value"], heading)
    fields = ["buyer", "supplier", "order_id", "invoice_id", "invoice_date", "payment_days_v1", "due_date_v1"]
    for row, key in enumerate(fields, 1):
        terms.write_row(row, 0, [LABELS[key], ledger[key]])
    terms.write_row(9, 0, ["Scope", "Fictional scenario; no real transactions"])
    for row, reference in enumerate(references, 11):
        terms.write_row(row, 0, ["Related record", reference])
    wb.close()


def render_docx(path: Path, ir: dict):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "MERIDIAN / WORKING RECORDS"
    section.footer.paragraphs[0].text = f"Fictional business record | {ir['version']}"
    section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.font.color.rgb = RGBColor.from_string("203549")
    doc.add_heading(ir["title"], 0)
    doc.add_paragraph(f"{ir['version']} | {ir['validity']}")
    for text in ir["paragraphs"]:
        doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Field", "Value"
    for item in ir["facts"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = item["label"], item["value"]
    if ir["lines"]:
        doc.add_heading("Transaction matching", 1)
        t = doc.add_table(rows=1, cols=5)
        t.style = "Light Shading Accent 1"
        for cell, value in zip(t.rows[0].cells, ["Product", "Ordered", "Received", "Unit USD", "Payable USD"]):
            cell.text = value
        for line in ir["lines"]:
            for cell, value in zip(t.add_row().cells, [line["product"], line["ordered"], line["received"], line["unit_price"], line["payable_amount"]]):
                cell.text = str(value)
    if ir["references"]:
        doc.add_heading("Related records", 1)
        for ref in ir["references"]:
            doc.add_paragraph(ref)
    doc.save(path)


def render_pptx(path: Path, ir: dict):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    def slide(title):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string("F4F7FA")
        box = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.8))
        p = box.text_frame.paragraphs[0]
        p.text, p.font.size, p.font.bold = title, Pt(30), True
        p.font.color.rgb = RGBColor.from_string("14334D")
        return s

    s = slide(ir["title"])
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(11.8), Inches(4.8))
    box.text_frame.word_wrap = True
    for index, text in enumerate([ir["validity"], *ir["paragraphs"]]):
        p = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
        p.text, p.font.size, p.space_after = text, Pt(20), Pt(18)
    s = slide("Settlement facts")
    facts = ir["facts"]
    table = s.shapes.add_table(len(facts)+1, 2, Inches(0.7), Inches(1.5), Inches(11.8), Inches(4.7)).table
    table.columns[0].width, table.columns[1].width = Inches(4.5), Inches(7.3)
    for row, values in enumerate([["Field", "Value"]] + [[f["label"], f["value"]] for f in facts]):
        for col, value in enumerate(values):
            cell = table.cell(row, col)
            cell.text = value
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(15)
    s = slide("Decision and supporting records")
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(11.8), Inches(4.5))
    box.text_frame.word_wrap = True
    box.text_frame.text = "Release the matched payable amount; retain unmatched quantities for supplier resolution.\n\n" + "\n".join(ir["references"])
    for p in box.text_frame.paragraphs:
        p.font.size = Pt(23)
    prs.save(path)


def render_pdf(path: Path, ir: dict):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(ir["title"]), styles["Title"]), Paragraph(escape(ir["validity"]), styles["Normal"]), Spacer(1, 14)]
    story += [Paragraph(escape(p), styles["BodyText"]) for p in ir["paragraphs"]]
    story += [Spacer(1, 14)]
    data = [["Field", "Value"]] + [[Paragraph(escape(f["label"]), styles["Normal"]), Paragraph(escape(f["value"]), styles["Normal"])] for f in ir["facts"]]
    table = Table(data, colWidths=[190, 305], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14334D")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white])]))
    story += [table, Spacer(1, 14)]
    story += [Paragraph(escape("Related record: " + ref), styles["BodyText"]) for ref in ir["references"]]
    SimpleDocTemplate(str(path), title=ir["title"], author="Fictional Meridian", leftMargin=48, rightMargin=48).build(story)


def generate(plan: FilePlan, world: WorldInstance, plans: list[FilePlan], workspace: Path, narrative=None) -> list[Path]:
    public = workspace / "public"
    target = safe_path(public, plan.relative_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    ir = content_ir(plan, world, plans, public, narrative) if plan.format != "xlsx" else {"title": "Settlement data", "facts": [], "derivation": None}
    fd, temp_name = tempfile.mkstemp(suffix=f".{plan.format}", dir=target.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        if plan.format == "xlsx":
            references = [next(p.relative_path for p in plans if p.id == d.file_id) for d in plan.dependencies]
            render_xlsx(temp, world, references)
        elif plan.format == "docx":
            render_docx(temp, ir)
        elif plan.format == "pptx":
            render_pptx(temp, ir)
        elif plan.format == "pdf":
            render_pdf(temp, ir)
        else:
            text = ir["title"] + "\n\n" + ir["validity"] + "\n\n" + "\n\n".join(ir["paragraphs"]) + "\n\n" + "\n".join(f"{f['label']}: {f['value']}" for f in ir["facts"]) + "\n\nRelated: " + ", ".join(ir["references"])
            if plan.format == "eml":
                message = EmailMessage(policy=SMTP)
                message["From"], message["To"] = "finance@example.invalid", "supplier@example.invalid"
                message["Subject"] = ir["title"]
                message["Date"] = format_datetime(datetime.fromisoformat(plan.business_date).replace(hour=12, tzinfo=timezone.utc))
                message["Message-ID"] = f"<{world.id}.{plan.id}@example.invalid>"
                message.set_content(text)
                temp.write_bytes(message.as_bytes())
            else:
                temp.write_text(text)
        os.replace(temp, target)
    finally:
        temp.unlink(missing_ok=True)
    ir_path = workspace / "private" / "ir" / f"{plan.id}.json"
    write_json(ir_path, ir)
    return [target, ir_path]
