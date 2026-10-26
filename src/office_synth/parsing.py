from __future__ import annotations

import re
import zipfile
from pathlib import Path, PurePosixPath

from .storage import digest

MIMES = {"html": "text/html", "pdf": "application/pdf", "md": "text/markdown", "txt": "text/plain",
         "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}


def detect(data: bytes, declared="", name="") -> str:
    head = data[:2048].lstrip().lower()
    if head.startswith(b"%pdf-"):
        return "pdf"
    if data.startswith(b"PK"):
        import io
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if any(PurePosixPath(n).is_absolute() or ".." in PurePosixPath(n).parts or "\\" in n for n in names):
                raise ValueError("Unsafe ZIP member path")
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024 or len(names) > 10000:
                raise ValueError("Expanded archive exceeds safety limits")
            if any("vbaProject" in n for n in names):
                raise ValueError("Macro-enabled sources are not supported")
            for prefix, fmt in [("word/", "docx"), ("xl/", "xlsx"), ("ppt/", "pptx")]:
                if any(n.startswith(prefix) for n in names):
                    return fmt
        raise ValueError("Unsupported ZIP archive")
    if any(tag in head for tag in [b"<!doctype html", b"<html", b"<body", b"<head"]):
        return "html"
    if data.startswith(bytes.fromhex("D0CF11E0")):
        raise ValueError("Legacy Office format is unsupported; convert explicitly")
    if b"\x00" not in head and ("text/" in declared or Path(name).suffix.lower() in {".txt", ".md", ".eml"}):
        data.decode("utf-8-sig")
        return "md" if name.lower().endswith(".md") else "txt"
    raise ValueError("Unsupported or unrecognized MIME")


def parse(path: Path, fmt: str) -> dict:
    data = path.read_bytes()
    units = []

    def add(text, **location):
        text = str(text).strip()
        if text:
            units.append({"text": text, "location": location})

    if fmt == "html":
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(data, "html.parser")
        for item in soup(["script", "style", "nav", "footer", "form"]):
            item.decompose()
        main = soup.find("main") or soup.find("article") or soup
        offset = 0
        for index, element in enumerate(main.find_all(["h1", "h2", "h3", "p", "li", "tr"])):
            text = element.get_text(" ", strip=True)
            add(text, node_index=index, tag=element.name, text_offset=offset)
            offset += len(text) + 1
        if not units:
            add(main.get_text(" ", strip=True), node_index=0, tag="document", text_offset=0)
    elif fmt in {"md", "txt", "eml"}:
        text = data.decode("utf-8-sig")
        for match in re.finditer(r"[^\n]+", text):
            add(match.group(), text_offset=match.start(), text_end=match.end())
    elif fmt == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF is not supported")
        for page, value in enumerate(reader.pages, 1):
            add(value.extract_text() or "", page=page)
    elif fmt == "docx":
        from docx import Document
        doc = Document(path)
        for index, p in enumerate(doc.paragraphs):
            add(p.text, paragraph=index)
        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    add(cell.text, table=ti, row=ri, column=ci)
        for si, section in enumerate(doc.sections):
            for kind in ["header", "footer"]:
                for pi, p in enumerate(getattr(section, kind).paragraphs):
                    add(p.text, section=si, area=kind, paragraph=pi)
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in wb:
                if sheet.max_row > 100000 or sheet.max_column > 1000:
                    raise ValueError("Worksheet dimensions exceed parser limits")
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            add(cell.value, sheet=sheet.title, range=cell.coordinate)
        finally:
            wb.close()
    elif fmt == "pptx":
        from pptx import Presentation
        for si, slide in enumerate(Presentation(path).slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    add(shape.text, slide=si, shape=shape.shape_id)
                elif shape.has_table:
                    for ri, row in enumerate(shape.table.rows):
                        for ci, cell in enumerate(row.cells):
                            add(cell.text, slide=si, shape=shape.shape_id, row=ri, column=ci)
    else:
        raise ValueError(f"Unsupported format: {fmt}")
    if not units:
        raise ValueError("No extractable text; scanned documents require an external OCR step")
    text = "\n".join(u["text"] for u in units)
    return {"format": fmt, "raw_hash": digest(data), "content_hash": digest(re.sub(r"\s+", " ", text).strip().casefold()), "units": units, "text": text}
